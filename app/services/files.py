"""File storage, kind detection and text extraction."""
from __future__ import annotations

import hashlib
import io
from pathlib import Path

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".tiff", ".ico"}
AUDIO_EXTS = {".wav", ".mp3", ".m4a", ".mp4", ".webm", ".ogg", ".oga", ".aac", ".flac"}
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml",
    ".html", ".htm", ".log", ".ini", ".cfg", ".conf", ".py", ".js", ".mjs", ".ts",
    ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".go", ".rs", ".rb",
    ".php", ".sh", ".bash", ".sql", ".css", ".scss", ".env", ".toml", ".r", ".kt",
}


def detect_kind(name: str, mime: str) -> str:
    ext = Path(name).suffix.lower()
    mime = (mime or "").lower()
    if ext == ".pdf" or mime == "application/pdf":
        return "pdf"
    if ext in AUDIO_EXTS or mime.startswith("audio/"):
        return "audio"
    if ext in IMAGE_EXTS or mime.startswith("image/"):
        return "image"
    if ext in {".doc", ".docx"}:
        return "doc"
    if ext in {".xls", ".xlsx", ".csv", ".tsv"}:
        return "sheet"
    if ext in {".ppt", ".pptx"}:
        return "slide"
    if ext in TEXT_EXTS or mime.startswith("text/"):
        return "text"
    return "other"


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def image_to_data_url(path: str, mime: str = "", max_side: int = 2560) -> str | None:
    """Downscale for vision models without building a second full-resolution raster."""
    try:
        import base64
        import io

        from PIL import Image, ImageOps

        with Image.open(path) as im:
            width, height = im.size
            if width <= 0 or height <= 0:
                return None
            # In-place thumbnailing can use decoder-level JPEG reduction where available.
            im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
            im = ImageOps.exif_transpose(im)
            if im.mode != "RGB":
                im = im.convert("RGB")
            buf = io.BytesIO()
            # Avoid chroma subsampling to preserve text and thin lines in screenshots.
            im.save(buf, format="JPEG", quality=90, subsampling=0)
            b64 = base64.b64encode(buf.getvalue()).decode()
        return f"data:image/jpeg;base64,{b64}"
    except Exception:
        pass
    try:
        import base64

        with open(path, "rb") as fh:
            b64 = base64.b64encode(fh.read()).decode()
        return f"data:{mime or 'image/png'};base64,{b64}"
    except Exception:
        return None

def decode_text(data: bytes) -> str:
    try:
        import chardet

        guess = chardet.detect(data[:30000]) or {}
        enc = guess.get("encoding") or "utf-8"
        return data.decode(enc, errors="replace")
    except Exception:
        return data.decode("utf-8", errors="replace")


def extract_text(kind: str, name: str, data: bytes) -> str:
    """Return plain text extracted from the file (best effort)."""
    try:
        if kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            parts = [(page.extract_text() or "") for page in reader.pages]
            return "\n\n".join(p for p in parts if p).strip()

        if kind == "doc":
            import docx

            doc = docx.Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines).strip()

        if kind == "sheet":
            if Path(name).suffix.lower() in {".csv", ".tsv"}:
                return decode_text(data)[:200000]
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"# Лист: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines).strip()

        if kind == "slide":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            lines = []
            for i, slide in enumerate(prs.slides, 1):
                lines.append(f"# Слайд {i}")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text:
                        lines.append(shape.text_frame.text)
            return "\n".join(lines).strip()

        if kind == "text":
            return decode_text(data).strip()
    except Exception as e:  # pragma: no cover - defensive
        return f"[Не удалось извлечь текст: {type(e).__name__}: {e}]"
    return ""
