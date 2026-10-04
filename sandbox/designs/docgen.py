"""docgen — real document generation for ChatStudio.

Produces genuinely presentable DOCX / PPTX / XLSX / PDF plus PNG chart images and
slide previews, driven by the style palettes in :mod:`designs.theme`.

Everything here is deterministic: callers pass a plain dict "spec" and get files
back.  No network access is required.
"""
from __future__ import annotations

import json
import math
import os
from pathlib import Path

from . import theme

# --------------------------------------------------------------------------- fonts
FONT_DIR = Path("/srv/fonts")
_REGISTERED = False

# matplotlib-safe families that ship with the sandbox (both cover Cyrillic)
MPL_SANS = "DejaVu Sans"
MPL_SERIF = "DejaVu Serif"


def register_fonts() -> list[str]:
    """Expose the bundled fonts (Yandex Sans, Lora, PT Serif) to matplotlib."""
    global _REGISTERED
    found: list[str] = []
    if _REGISTERED:
        return found
    try:
        import matplotlib
        from matplotlib import font_manager as fm

        cache = Path("/tmp/mplcache")
        cache.mkdir(parents=True, exist_ok=True)
        matplotlib.use("Agg")
        matplotlib.rcParams["font.family"] = MPL_SANS
        paths = []
        if FONT_DIR.is_dir():
            paths = [str(p) for p in FONT_DIR.glob("*.ttf")] + [
                str(p) for p in FONT_DIR.glob("*.otf")
            ]
        for p in paths:
            try:
                fm.fontManager.addfont(p)
                found.append(Path(p).stem)
            except Exception:
                continue
        if found:
            names = sorted({fm.FontProperties(fname=p).get_name() for p in paths})
            matplotlib.rcParams["font.sans-serif"] = names + [MPL_SANS]
            matplotlib.rcParams["font.family"] = "sans-serif"
    except Exception:
        pass
    _REGISTERED = True
    return found


def mpl_font(prefer_serif: bool = False) -> str:
    register_fonts()
    try:
        import matplotlib
        from matplotlib import font_manager as fm

        want = ["Lora", "PT Serif"] if prefer_serif else ["YS Text", "Inter", "Lora"]
        have = {f.name for f in fm.fontManager.ttflist}
        for w in want:
            if w in have:
                return w
    except Exception:
        pass
    return MPL_SERIF if prefer_serif else MPL_SANS


# --------------------------------------------------------------------------- helpers
def hx(color: str) -> str:
    return color if str(color).startswith("#") else "#" + str(color)


def _mix(c1: str, c2: str, t: float) -> str:
    t = max(0.0, min(1.0, t))
    a = [int(hx(c1)[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(hx(c2)[i:i + 2], 16) for i in (1, 3, 5)]
    return "#%02X%02X%02X" % tuple(round(a[i] + (b[i] - a[i]) * t) for i in range(3))


def _contrast(bg: str) -> str:
    r, g, b = (int(hx(bg)[i:i + 2], 16) for i in (1, 3, 5))
    return "#111111" if (0.299 * r + 0.587 * g + 0.114 * b) > 150 else "#FFFFFF"


def ensure_spec(spec) -> dict:
    if isinstance(spec, str):
        try:
            spec = json.loads(spec)
        except Exception:
            spec = {"title": spec, "blocks": []}
    spec = spec or {}
    spec.setdefault("title", "Документ")
    spec.setdefault("subtitle", "")
    spec.setdefault("author", "")
    spec.setdefault("style", "minimal")
    spec.setdefault("blocks", [])
    return spec


# =========================================================================== CHARTS
CHART_KINDS = ("bar", "hbar", "line", "area", "pie", "donut", "scatter", "stacked")


def chart(path: str | Path, spec: dict, style: str = "minimal") -> str:
    """Draw a chart PNG.  ``spec`` = {kind, title, labels, series:[{name,values}], xlabel, ylabel}."""
    import matplotlib
    matplotlib.use("Agg")
    register_fonts()
    import matplotlib.pyplot as plt

    p = theme.palette(style)
    kind = (spec.get("kind") or "bar").lower()
    labels = spec.get("labels") or []
    series = spec.get("series") or [{"name": spec.get("series_name", "Значение"),
                                     "values": spec.get("values") or []}]
    cols = theme.colors(style, max(3, len(series) + 1))

    fig, ax = plt.subplots(figsize=spec.get("figsize", (9.6, 5.4)), dpi=160)
    fig.patch.set_facecolor(hx(p["bg"]))
    ax.set_facecolor(hx(p["bg"]))
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(_mix(p["muted"], p["bg"], 0.55))
        ax.spines[side].set_linewidth(0.9)
    ax.tick_params(colors=hx(p["muted"]), labelsize=10)
    ax.grid(axis="y", color=_mix(p["muted"], p["bg"], 0.82), linewidth=0.8)
    ax.set_axisbelow(True)

    fname = mpl_font()
    title = spec.get("title") or ""
    if title:
        ax.set_title(title, fontsize=15, fontweight="bold", color=hx(p["text"]),
                     loc="left", pad=14, fontfamily=fname)
    if spec.get("xlabel"):
        ax.set_xlabel(spec["xlabel"], color=hx(p["muted"]), fontsize=10)
    if spec.get("ylabel"):
        ax.set_ylabel(spec["ylabel"], color=hx(p["muted"]), fontsize=10)

    xs = range(len(labels)) if labels else None
    if kind in ("bar", "stacked"):
        width = 0.8 / max(1, len(series))
        bottom = [0.0] * (len(labels) or len(series[0]["values"]))
        for i, s in enumerate(series):
            vals = s.get("values") or []
            pos = [x + i * width - 0.4 + width / 2 for x in (xs or range(len(vals)))]
            if kind == "stacked":
                ax.bar(pos, vals, width, bottom=bottom, label=s.get("name", ""),
                       color=hx(cols[i % len(cols)]), edgecolor="none")
                bottom = [bottom[j] + (vals[j] if j < len(vals) else 0) for j in range(len(bottom))]
            else:
                ax.bar(pos, vals, width * 0.92, label=s.get("name", ""),
                       color=hx(cols[i % len(cols)]), edgecolor="none")
        if labels:
            ax.set_xticks([x for x in xs])
            ax.set_xticklabels(labels, color=hx(p["muted"]), fontsize=10)
    elif kind == "hbar":
        vals = series[0]["values"]
        pos = list(range(len(vals)))
        ax.barh(pos, vals, color=[hx(c) for c in cols[:len(vals)]] or hx(p["primary"]), height=0.62)
        ax.set_yticks(pos)
        ax.set_yticklabels(labels or [str(i + 1) for i in pos], color=hx(p["muted"]), fontsize=10)
        ax.invert_yaxis()
        ax.grid(axis="y", visible=False)
        ax.grid(axis="x", color=_mix(p["muted"], p["bg"], 0.82), linewidth=0.8)
    elif kind in ("line", "area"):
        for i, s in enumerate(series):
            vals = s.get("values") or []
            x = list(xs) if labels and len(labels) == len(vals) else list(range(len(vals)))
            if kind == "area":
                ax.fill_between(x, vals, color=hx(cols[i % len(cols)]), alpha=0.22)
            ax.plot(x, vals, color=hx(cols[i % len(cols)]), linewidth=2.4,
                    marker="o", markersize=4.5, label=s.get("name", ""))
        if labels:
            ax.set_xticks(list(xs))
            ax.set_xticklabels(labels, color=hx(p["muted"]), fontsize=10)
    elif kind in ("pie", "donut"):
        vals = series[0]["values"]
        wedge = {"width": 0.42} if kind == "donut" else {}
        wedges, texts, autos = ax.pie(
            vals, labels=labels or None, autopct="%1.0f%%", startangle=90,
            colors=[hx(c) for c in (cols * 3)[:len(vals)]], wedgeprops=wedge,
            textprops={"color": hx(p["text"]), "fontsize": 10, "fontfamily": fname},
        )
        for a in autos:
            a.set_color(_contrast(p["bg"]))
            a.set_fontsize(9.5)
        ax.grid(False)
        ax.set_aspect("equal")
    elif kind == "scatter":
        for i, s in enumerate(series):
            ax.scatter(s.get("x") or list(range(len(s.get("values") or []))),
                       s.get("values") or [], s=46, color=hx(cols[i % len(cols)]),
                       label=s.get("name", ""), alpha=0.85, edgecolors=hx(p["bg"]), linewidths=1)
    else:
        ax.bar(list(xs or range(len(series[0]["values"]))), series[0]["values"], color=hx(p["primary"]))

    if len(series) > 1 and kind not in ("pie", "donut"):
        lg = ax.legend(frameon=False, fontsize=10, labelcolor=hx(p["text"]))
        lg.get_frame().set_alpha(0)
    fig.tight_layout()
    out = str(path)
    fig.savefig(out, facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)
    return out


# =========================================================================== PPTX
SLIDE_W = 13.333
SLIDE_H = 7.5


def _pptx_setup(spec: dict, style: str):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def _rect(slide, x, y, w, h, fill=None, line=None, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = _rgb(line)
        sh.line.width = _emu(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _rgb(color: str):
    from pptx.dml.color import RGBColor
    c = hx(color).lstrip("#")
    return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))


def _emu(points: float):
    from pptx.util import Pt
    return Pt(points)


def _text(slide, text, x, y, w, h, size=18, color="#222", bold=False, align="left",
          font=None, italic=False, spacing=1.0, anchor="top"):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                          "right": PP_ALIGN.RIGHT}[align]
        para.line_spacing = spacing
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = _rgb(color)
        if font:
            run.font.name = font
    return box


def _footer(slide, p, page: int, total: int, font: str, label: str = ""):
    _rect(slide, 0.7, SLIDE_H - 0.62, SLIDE_W - 1.4, 0.012, fill=_mix(p["muted"], p["bg"], 0.72))
    if label:
        _text(slide, label, 0.7, SLIDE_H - 0.52, 8.0, 0.3, size=9.5, color=p["muted"], font=font)
    _text(slide, f"{page} / {total}", SLIDE_W - 2.0, SLIDE_H - 0.52, 1.3, 0.3,
          size=9.5, color=p["muted"], align="right", font=font)


def _slide_header(slide, p, title: str, font: str, kicker: str = ""):
    _rect(slide, 0.7, 0.62, 0.09, 0.62, fill=p["primary"])
    y = 0.6
    if kicker:
        _text(slide, kicker.upper(), 0.98, 0.52, 9.0, 0.28, size=10.5,
              color=p["accent"], bold=True, font=font)
        y = 0.82
    _text(slide, title, 0.98, y, SLIDE_W - 2.0, 0.8, size=27, color=p["text"],
          bold=True, font=font, spacing=0.95)


def pptx_deck(path: str | Path, spec: dict, style: str | None = None) -> str:
    """Build a presentation.  spec = {title, subtitle, author, blocks:[...]}.

    Block kinds: bullets | two_col | metrics | chart | image | quote | section | text | table
    """
    from pptx.util import Inches

    spec = ensure_spec(spec)
    style = style or spec["style"]
    p = theme.palette(style)
    font = p.get("font") or "Segoe UI"
    prs = _pptx_setup(spec, style)
    blank = prs.slide_layouts[6]

    blocks = [b for b in spec["blocks"] if isinstance(b, dict)]
    total = 1 + len(blocks) + (1 if spec.get("thanks", True) else 0)
    page = 0

    # ---- title slide
    s = prs.slides.add_slide(blank)
    page += 1
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=p["bg"])
    _rect(s, 0, 0, SLIDE_W, 0.34, fill=p["primary"])
    _rect(s, 0, SLIDE_H - 0.9, SLIDE_W, 0.9, fill=_mix(p["primary"], p["bg"], 0.9))
    _rect(s, 1.1, 2.05, 0.16, 1.5, fill=p["accent"])
    _text(s, spec["title"], 1.5, 2.0, SLIDE_W - 3.0, 1.6, size=42, color=p["text"],
          bold=True, font=font, spacing=0.94)
    if spec.get("subtitle"):
        _text(s, spec["subtitle"], 1.5, 3.75, SLIDE_W - 3.0, 0.9, size=17,
              color=p["muted"], font=font, spacing=1.15)
    meta = " · ".join(x for x in [spec.get("author"), spec.get("date")] if x)
    if meta:
        _text(s, meta, 1.5, 6.0, SLIDE_W - 3.0, 0.4, size=12, color=p["muted"], font=font)
    _text(s, style, SLIDE_W - 3.2, 6.0, 2.5, 0.4, size=11, color=p["accent"],
          align="right", font=font)

    # ---- content slides
    for b in blocks:
        kind = (b.get("kind") or "bullets").lower()
        s = prs.slides.add_slide(blank)
        page += 1
        _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=p["bg"])
        kicker = b.get("kicker", "")
        title = b.get("title", "")

        if kind == "section":
            _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=p["primary"])
            _text(s, str(b.get("index", "")), 1.1, 2.25, 2.0, 1.2, size=64,
                  color=_mix(p["primary"], "#FFFFFF", 0.45), bold=True, font=font)
            _text(s, title, 1.1, 3.3, SLIDE_W - 2.2, 1.2, size=34, color=_contrast(p["primary"]),
                  bold=True, font=font)
            if b.get("text"):
                _text(s, b["text"], 1.1, 4.5, SLIDE_W - 2.2, 0.8, size=15,
                      color=_mix(p["primary"], "#FFFFFF", 0.72), font=font)
            continue

        if kind == "quote":
            _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=_mix(p["bg"], p["surface"], 0.7))
            _rect(s, 1.6, 2.35, 0.13, 2.3, fill=p["accent"])
            _text(s, "«" + (b.get("text") or "") + "»", 2.1, 2.4, SLIDE_W - 4.2, 2.0,
                  size=25, color=p["text"], font=font, spacing=1.2, italic=True)
            if b.get("author"):
                _text(s, "— " + b["author"], 2.1, 4.7, SLIDE_W - 4.2, 0.5, size=14,
                      color=p["muted"], font=font)
            _footer(s, p, page, total, font, spec["title"])
            continue

        _slide_header(s, p, title, font, kicker)

        if kind == "bullets":
            items = b.get("items") or []
            top = 1.75
            room = SLIDE_H - top - 1.0
            step = min(0.72, room / max(1, len(items)))
            for i, it in enumerate(items):
                y = top + i * step
                _rect(s, 0.98, y + 0.12, 0.085, 0.085, fill=p["accent"], radius=0.5)
                txt = it if isinstance(it, str) else str(it)
                size = 15.5
                _text(s, txt, 1.28, y, SLIDE_W - 2.3, step, size=size, color=p["text"],
                      font=font, spacing=1.18, anchor="top")
        elif kind == "two_col":
            colw = (SLIDE_W - 2.1) / 2 - 0.2
            for idx, key in enumerate(("left", "right")):
                col = b.get(key) or {}
                x = 0.98 + idx * (colw + 0.42)
                if col.get("title"):
                    _text(s, col["title"], x, 1.7, colw, 0.4, size=16, bold=True,
                          color=p["primary"], font=font)
                _rect(s, x, 2.15, colw, 0.01, fill=_mix(p["muted"], p["bg"], 0.7))
                _text(s, col.get("text", ""), x, 2.3, colw, SLIDE_H - 3.3, size=13,
                      color=p["text"], font=font, spacing=1.22)
        elif kind == "metrics":
            cards = b.get("items") or []
            n = max(1, len(cards))
            gap = 0.35
            cw = (SLIDE_W - 1.6 - gap * (n - 1)) / n
            for i, c in enumerate(cards):
                x = 0.8 + i * (cw + gap)
                _rect(s, x, 2.3, cw, 2.35, fill=_mix(p["surface"], p["bg"], 0.15),
                      line=_mix(p["muted"], p["bg"], 0.7), radius=0.06)
                _rect(s, x, 2.3, cw, 0.075, fill=p["accent"] if i % 2 else p["primary"])
                _text(s, str(c.get("value", "")), x + 0.22, 2.72, cw - 0.44, 0.95,
                      size=34, bold=True, color=p["primary"], font=font)
                _text(s, str(c.get("label", "")), x + 0.22, 3.72, cw - 0.44, 0.75,
                      size=12.5, color=p["muted"], font=font, spacing=1.15)
        elif kind == "chart":
            img = b.get("image")
            if img and os.path.exists(img):
                from PIL import Image
                with Image.open(img) as im:
                    ratio = im.width / im.height
                w = min(SLIDE_W - 2.6, 9.6)
                h = w / ratio
                if h > SLIDE_H - 2.5:
                    h = SLIDE_H - 2.5
                    w = h * ratio
                s.shapes.add_picture(img, Inches((SLIDE_W - w) / 2), Inches(1.85),
                                     inches_w(w), inches_h(h))
            if b.get("note"):
                _text(s, b["note"], 0.98, SLIDE_H - 1.15, SLIDE_W - 2.0, 0.4, size=11.5,
                      color=p["muted"], font=font)
        elif kind == "image":
            img = b.get("image")
            if img and os.path.exists(img):
                from PIL import Image
                with Image.open(img) as im:
                    ratio = im.width / im.height
                w = min(SLIDE_W - 3.0, 8.6)
                h = w / ratio
                if h > SLIDE_H - 2.6:
                    h = SLIDE_H - 2.6
                    w = h * ratio
                s.shapes.add_picture(img, Inches((SLIDE_W - w) / 2), Inches(1.8),
                                     inches_w(w), inches_h(h))
            if b.get("caption"):
                _text(s, b["caption"], 0.98, SLIDE_H - 1.1, SLIDE_W - 2.0, 0.4,
                      size=11.5, color=p["muted"], align="center", font=font)
        elif kind == "table":
            rows = (b.get("rows") or [])[:8]
            head = b.get("head") or []
            cols = max([len(head)] + [len(r) for r in rows] or [1])
            n = len(rows) + (1 if head else 0)
            shape = s.shapes.add_table(n, cols, Inches(0.9), Inches(1.8),
                                       Inches(SLIDE_W - 1.8), Inches(0.42 * n))
            tbl = shape.table
            for ci, htxt in enumerate(head):
                cell = tbl.cell(0, ci)
                cell.text = str(htxt)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(p["primary"])
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = _emu(12)
                        run.font.bold = True
                        run.font.name = font
                        run.font.color.rgb = _rgb(_contrast(p["primary"]))
            for ri, row in enumerate(rows, start=1 if head else 0):
                for ci in range(cols):
                    cell = tbl.cell(ri, ci)
                    cell.text = str(row[ci]) if ci < len(row) else ""
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(
                        _mix(p["surface"], p["bg"], 0.15) if ri % 2 else p["bg"])
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = _emu(11)
                            run.font.name = font
                            run.font.color.rgb = _rgb(p["text"])
        else:  # plain text
            _text(s, b.get("text", ""), 0.98, 1.9, SLIDE_W - 2.0, SLIDE_H - 3.3,
                  size=15, color=p["text"], font=font, spacing=1.25)

        _footer(s, p, page, total, font, spec["title"])

    # ---- closing
    if spec.get("thanks", True):
        s = prs.slides.add_slide(blank)
        page += 1
        _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=p["primary"])
        _text(s, spec.get("closing", "Спасибо!"), 1.1, 2.6, SLIDE_W - 2.2, 1.1,
              size=40, bold=True, color=_contrast(p["primary"]), font=font)
        if spec.get("author"):
            _text(s, spec["author"], 1.1, 3.9, SLIDE_W - 2.2, 0.5, size=15,
                  color=_mix(p["primary"], "#FFFFFF", 0.7), font=font)

    out = str(path)
    prs.save(out)
    return out


def inches_h(v):
    from pptx.util import Inches
    return Inches(v)


def inches_w(v):
    from pptx.util import Inches
    return Inches(v)


# =========================================================================== DOCX
def _docx_page_number(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for el, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        node = OxmlElement(el)
        for k, v in attrs.items():
            node.set(qn(k), v)
        if text:
            node.text = text
        run._r.append(node)
    return run


def _docx_shade(cell, color: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hx(color).lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def _docx_no_borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        borders.append(el)
    pr.append(borders)


def docx_report(path: str | Path, spec: dict, style: str | None = None) -> str:
    """Build a formatted Word document with a cover page, headings, tables and page numbers."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    spec = ensure_spec(spec)
    style = style or spec["style"]
    p = theme.palette(style)
    font = p.get("font") or "Calibri"

    doc = Document()
    theme.docx_style(doc, style)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(72)
        section.left_margin = section.right_margin = Pt(72)

    def para(text="", size=11, color=None, bold=False, italic=False, align=None,
             space_after=7, spacing=1.28):
        par = doc.add_paragraph()
        if align:
            par.alignment = align
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.line_spacing = spacing
        run = par.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        try:
            run.font.color.rgb = RGBColor.from_string(hx(color or p["text"]).lstrip("#"))
        except Exception:
            pass
        return par

    # ---- cover
    if spec.get("cover", True):
        para("", space_after=90)
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(" " * 1)
        r.font.size = Pt(2)
        bar = doc.add_table(rows=1, cols=1)
        cover = bar.cell(0, 0)
        cover.text = ""
        _docx_shade(cover, p["primary"])
        cover.paragraphs[0].add_run(" ").font.size = Pt(8)
        para("", space_after=14)
        para(spec["title"], size=30, bold=True, color=p["text"], space_after=8, spacing=1.1)
        if spec.get("subtitle"):
            para(spec["subtitle"], size=14, color=p["muted"], space_after=20)
        meta = " · ".join(x for x in [spec.get("author"), spec.get("date")] if x)
        if meta:
            para(meta, size=11, color=p["muted"])
        doc.add_page_break()

    # ---- body
    for b in spec["blocks"]:
        if not isinstance(b, dict):
            para(str(b))
            continue
        kind = (b.get("kind") or "text").lower()
        if kind == "h1":
            h = doc.add_heading(str(b.get("text", "")), level=1)
            h.paragraph_format.space_before = Pt(16)
        elif kind == "h2":
            h = doc.add_heading(str(b.get("text", "")), level=2)
            h.paragraph_format.space_before = Pt(12)
        elif kind == "h3":
            doc.add_heading(str(b.get("text", "")), level=3)
        elif kind == "bullets":
            for it in b.get("items") or []:
                par = doc.add_paragraph(style="List Bullet")
                par.paragraph_format.space_after = Pt(4)
                run = par.add_run(str(it))
                run.font.size = Pt(11)
                run.font.name = font
        elif kind == "numbers":
            for it in b.get("items") or []:
                par = doc.add_paragraph(style="List Number")
                par.paragraph_format.space_after = Pt(4)
                run = par.add_run(str(it))
                run.font.size = Pt(11)
                run.font.name = font
        elif kind == "callout":
            t = doc.add_table(rows=1, cols=1)
            cell = t.cell(0, 0)
            _docx_shade(cell, _mix(p["accent"], "#FFFFFF", 0.84))
            cell.text = ""
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(6)
            cp.paragraph_format.space_after = Pt(6)
            run = cp.add_run(str(b.get("text", "")))
            run.font.size = Pt(11)
            run.font.name = font
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(hx(p["primary"]).lstrip("#"))
        elif kind == "table":
            head = b.get("head") or []
            rows = b.get("rows") or []
            cols = max([len(head)] + [len(r) for r in rows] or [1])
            t = doc.add_table(rows=len(rows) + (1 if head else 0), cols=cols)
            t.style = "Table Grid"
            off = 0
            if head:
                for i, htxt in enumerate(head):
                    cell = t.cell(0, i)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(str(htxt))
                    run.font.bold = True
                    run.font.size = Pt(10.5)
                    run.font.name = font
                    run.font.color.rgb = RGBColor.from_string(_contrast(p["primary"]).lstrip("#"))
                    _docx_shade(cell, p["primary"])
                off = 1
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    cell = t.cell(ri + off, ci)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(str(row[ci]) if ci < len(row) else "")
                    run.font.size = Pt(10.5)
                    run.font.name = font
                    if ri % 2:
                        _docx_shade(cell, _mix(p["surface"], p["bg"], 0.2))
            para("", space_after=6)
        elif kind == "image" and b.get("image") and os.path.exists(b["image"]):
            from docx.shared import Inches
            doc.add_picture(b["image"], width=Inches(6.1))
            if b.get("caption"):
                para(b["caption"], size=9.5, color=p["muted"],
                     align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            para(str(b.get("text", "")))

    # ---- footer with page numbers
    for section in doc.sections:
        footer = section.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = fp.add_run(spec["title"] + "   ")
        run.font.size = Pt(8.5)
        run.font.name = font
        try:
            run.font.color.rgb = RGBColor.from_string(hx(p["muted"]).lstrip("#"))
        except Exception:
            pass
        _docx_page_number(fp)

    out = str(path)
    doc.save(out)
    return out


# =========================================================================== XLSX
def xlsx_report(path: str | Path, spec: dict, style: str | None = None) -> str:
    """Styled workbook: header row, zebra, number formats, totals, freeze panes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    spec = ensure_spec(spec)
    style = style or spec["style"]
    p = theme.palette(style)
    font = p.get("font") or "Calibri"

    wb = Workbook()
    wb.remove(wb.active)
    sheets = spec.get("sheets") or [{
        "name": spec.get("title", "Лист1")[:28],
        "head": spec.get("head") or [],
        "rows": spec.get("rows") or [],
    }]

    thin = Side(style="thin", color=hx(_mix(p["muted"], p["bg"], 0.75)).lstrip("#"))
    for sh in sheets:
        ws = wb.create_sheet(title=str(sh.get("name") or "Лист")[:28])
        head = sh.get("head") or []
        rows = sh.get("rows") or []
        if head:
            ws.append(head)
        for r in rows:
            ws.append(list(r))
        ncols = max([len(head)] + [len(r) for r in rows] or [1])

        if head:
            for c in range(1, ncols + 1):
                cell = ws.cell(row=1, column=c)
                cell.fill = PatternFill("solid", fgColor=hx(p["primary"]).lstrip("#"))
                cell.font = Font(bold=True, color=hx(_contrast(p["primary"])).lstrip("#"),
                                 name=font, size=11)
                cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
            ws.row_dimensions[1].height = 26

        zebra = PatternFill("solid", fgColor=hx(_mix(p["surface"], p["bg"], 0.15)).lstrip("#"))
        first = 2 if head else 1
        for ridx in range(first, ws.max_row + 1):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=ridx, column=c)
                cell.font = Font(name=font, size=10.5,
                                 color=hx(p["text"]).lstrip("#"))
                cell.border = Border(bottom=thin)
                if (ridx - first) % 2:
                    cell.fill = zebra
                if isinstance(cell.value, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = sh.get("number_format", "#,##0.00")

        for c in range(1, ncols + 1):
            longest = 0
            for ridx in range(1, ws.max_row + 1):
                v = ws.cell(row=ridx, column=c).value
                longest = max(longest, len(str(v)) if v is not None else 0)
            ws.column_dimensions[get_column_letter(c)].width = min(46, max(11, longest + 3))

        if sh.get("totals") and rows:
            total_row = ws.max_row + 1
            ws.cell(row=total_row, column=1, value=sh.get("totals_label", "Итого"))
            for c in range(1, ncols + 1):
                cell = ws.cell(row=total_row, column=c)
                cell.font = Font(bold=True, name=font, color=hx(p["primary"]).lstrip("#"))
                cell.fill = PatternFill("solid", fgColor=hx(_mix(p["accent"], p["bg"], 0.82)).lstrip("#"))
                if c > 1:
                    col = get_column_letter(c)
                    cell.value = f"=SUM({col}{first}:{col}{ws.max_row - 1})"
                    cell.number_format = sh.get("number_format", "#,##0.00")
                    cell.alignment = Alignment(horizontal="right")

        ws.freeze_panes = "A2" if head else "A1"

    out = str(path)
    wb.save(out)
    return out


# =========================================================================== PDF (matplotlib)
def pdf_report(path: str | Path, spec: dict, style: str | None = None) -> str:
    """Multi-page PDF rendered with matplotlib (works without LibreOffice)."""
    import matplotlib
    matplotlib.use("Agg")
    register_fonts()
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    spec = ensure_spec(spec)
    style = style or spec["style"]
    p = theme.palette(style)
    font = mpl_font()
    out = str(path)

    A4 = (8.27, 11.69)  # inches
    with PdfPages(out) as pdf:
        # cover
        fig = plt.figure(figsize=A4)
        fig.patch.set_facecolor(hx(p["bg"]))
        fig.text(0.08, 0.90, spec["title"], fontsize=26, fontweight="bold",
                 color=hx(p["text"]), fontfamily=font, wrap=True)
        fig.patches.append(plt.Rectangle((0.08, 0.845), 0.16, 0.012,
                                         transform=fig.transFigure,
                                         facecolor=hx(p["accent"]), edgecolor="none"))
        if spec.get("subtitle"):
            fig.text(0.08, 0.80, spec["subtitle"], fontsize=13, color=hx(p["muted"]),
                     fontfamily=font, wrap=True)
        meta = " · ".join(x for x in [spec.get("author"), spec.get("date")] if x)
        if meta:
            fig.text(0.08, 0.10, meta, fontsize=10, color=hx(p["muted"]), fontfamily=font)
        pdf.savefig(fig, facecolor=fig.get_facecolor())
        plt.close(fig)

        # content pages
        buffer: list = []
        y = 0.92

        def flush():
            nonlocal buffer, y
            if not buffer:
                return
            fig = plt.figure(figsize=A4)
            fig.patch.set_facecolor(hx(p["bg"]))
            yy = y
            for kind, payload in buffer:
                if kind == "h1":
                    fig.text(0.08, yy, payload, fontsize=17, fontweight="bold",
                             color=hx(p["primary"]), fontfamily=font, wrap=True)
                    yy -= 0.045
                elif kind == "h2":
                    fig.text(0.08, yy, payload, fontsize=13.5, fontweight="bold",
                             color=hx(p["text"]), fontfamily=font, wrap=True)
                    yy -= 0.038
                elif kind == "bullet":
                    fig.text(0.10, yy, "•  " + payload, fontsize=11,
                             color=hx(p["text"]), fontfamily=font, wrap=True)
                    yy -= 0.032
                elif kind == "image":
                    try:
                        from PIL import Image
                        with Image.open(payload) as im:
                            ratio = im.height / im.width
                        h = min(0.42, 6.6 / 8.27 * ratio * 0.5)
                        ax = fig.add_axes([0.10, yy - h, 0.80, h])
                        ax.imshow(im)
                        ax.axis("off")
                        yy -= h + 0.03
                    except Exception:
                        pass
                else:
                    words = str(payload)
                    fig.text(0.08, yy, words, fontsize=11, color=hx(p["text"]),
                             fontfamily=font, wrap=True, linespacing=1.5)
                    yy -= 0.032 * (1 + len(words) // 95)
            pdf.savefig(fig, facecolor=fig.get_facecolor())
            plt.close(fig)
            buffer = []
            y = 0.92

        for b in spec["blocks"]:
            if not isinstance(b, dict):
                buffer.append(("text", str(b)))
                continue
            kind = (b.get("kind") or "text").lower()
            if kind in ("h1", "h2", "h3"):
                if len(buffer) > 16:
                    flush()
                buffer.append(("h1" if kind == "h1" else "h2", b.get("text", "")))
            elif kind in ("bullets", "numbers"):
                for it in b.get("items") or []:
                    buffer.append(("bullet", str(it)))
            elif kind == "image" and b.get("image") and os.path.exists(b["image"]):
                buffer.append(("image", b["image"]))
            elif kind == "pagebreak":
                flush()
            else:
                buffer.append(("text", b.get("text", "")))
            if len(buffer) > 26:
                flush()
        flush()

        d = pdf.infodict()
        d["Title"] = spec["title"]
        d["Author"] = spec.get("author", "ChatStudio")
    return out


# =========================================================================== previews
def slide_preview(path: str | Path, slide: dict, style: str = "minimal",
                  size=(1280, 720), index: int = 1, total: int = 1,
                  deck_title: str = "") -> str:
    """Render a PNG that mirrors how a PPTX slide from pptx_deck will look."""
    import matplotlib
    matplotlib.use("Agg")
    register_fonts()
    import matplotlib.pyplot as plt

    p = theme.palette(style)
    font = mpl_font()
    W, H = size
    fig = plt.figure(figsize=(W / 100, H / 100), dpi=100)
    fig.patch.set_facecolor(hx(p["bg"]))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    ax.set_facecolor(hx(p["bg"]))

    def band(x, y, w, h, color, alpha=1.0):
        ax.add_patch(plt.Rectangle((x, y), w, h, facecolor=hx(color),
                                   edgecolor="none", alpha=alpha))

    kind = (slide.get("kind") or "bullets").lower()

    if kind == "title":
        band(0, 97, 100, 3, p["primary"])
        band(0, 0, 100, 12, _mix(p["primary"], p["bg"], 0.9))
        band(8, 45, 1.2, 22, p["accent"])
        ax.text(11, 62, slide.get("title", ""), fontsize=30, fontweight="bold",
                color=hx(p["text"]), fontfamily=font, va="top", wrap=True)
        if slide.get("subtitle"):
            ax.text(11, 44, slide["subtitle"], fontsize=13, color=hx(p["muted"]),
                    fontfamily=font, va="top", wrap=True)
        meta = slide.get("author", "")
        if meta:
            ax.text(11, 11, meta, fontsize=10, color=hx(p["muted"]), fontfamily=font)
    elif kind == "section":
        band(0, 0, 100, 100, p["primary"])
        ax.text(8, 62, slide.get("title", ""), fontsize=28, fontweight="bold",
                color="#FFFFFF", fontfamily=font, va="top", wrap=True)
        if slide.get("text"):
            ax.text(8, 48, slide["text"], fontsize=12,
                    color=_mix(p["primary"], "#FFFFFF", 0.72), fontfamily=font, va="top")
    elif kind == "quote":
        band(0, 0, 100, 100, _mix(p["bg"], p["surface"], 0.7))
        band(12, 30, 1.0, 40, p["accent"])
        ax.text(17, 68, "«" + slide.get("text", "") + "»", fontsize=19, style="italic",
                color=hx(p["text"]), fontfamily=font, va="top", wrap=True)
        if slide.get("author"):
            ax.text(17, 32, "— " + slide["author"], fontsize=12, color=hx(p["muted"]),
                    fontfamily=font, va="top")
    elif kind == "metrics":
        cards = slide.get("items") or []
        n = max(1, len(cards))
        gap = 3
        cw = (100 - 12 - gap * (n - 1)) / n
        for i, c in enumerate(cards):
            x = 6 + i * (cw + gap)
            band(x, 32, cw, 34, _mix(p["surface"], p["bg"], 0.15))
            ax.add_patch(plt.Rectangle((x, 32), cw, 34, fill=False,
                                       edgecolor=hx(_mix(p["muted"], p["bg"], 0.7)), lw=1))
            band(x, 65.2, cw, 0.8, p["accent"] if i % 2 else p["primary"])
            ax.text(x + cw / 2, 56, str(c.get("value", "")), fontsize=24, fontweight="bold",
                    color=hx(p["primary"]), ha="center", fontfamily=font, va="center")
            ax.text(x + cw / 2, 40, str(c.get("label", "")), fontsize=10,
                    color=hx(p["muted"]), ha="center", fontfamily=font, va="center", wrap=True)
        band(5, 88, 0.8, 5, p["primary"])
        ax.text(7, 93, slide.get("title", ""), fontsize=18, fontweight="bold",
                color=hx(p["text"]), fontfamily=font, va="top", wrap=True)
    elif kind == "chart" and slide.get("image") and os.path.exists(slide["image"]):
        band(5, 88, 0.8, 5, p["primary"])
        ax.text(7, 93, slide.get("title", ""), fontsize=18, fontweight="bold",
                color=hx(p["text"]), fontfamily=font, va="top", wrap=True)
        try:
            from PIL import Image
            with Image.open(slide["image"]) as im:
                ax.imshow(im, extent=(14, 86, 16, 82), aspect="auto")
        except Exception:
            pass
    else:  # bullets / two_col / text
        band(5, 88, 0.8, 5, p["primary"])
        if slide.get("kicker"):
            ax.text(7, 95.5, str(slide["kicker"]).upper(), fontsize=8.5,
                    color=hx(p["accent"]), fontweight="bold", fontfamily=font, va="top")
        ax.text(7, 93, slide.get("title", ""), fontsize=18, fontweight="bold",
                color=hx(p["text"]), fontfamily=font, va="top", wrap=True)
        if kind == "two_col":
            for idx, key in enumerate(("left", "right")):
                col = slide.get(key) or {}
                x = 6 + idx * 46
                if col.get("title"):
                    ax.text(x, 82, col["title"], fontsize=12, fontweight="bold",
                            color=hx(p["primary"]), fontfamily=font, va="top")
                ax.text(x, 78, col.get("text", ""), fontsize=10, color=hx(p["text"]),
                        fontfamily=font, va="top", wrap=True)
        else:
            items = slide.get("items") or [slide.get("text", "")]
            y = 82
            step = min(11, 62 / max(1, len(items)))
            for it in items:
                ax.add_patch(plt.Circle((7.4, y - 1.2), 0.55, color=hx(p["accent"])))
                ax.text(9.5, y, str(it), fontsize=11.5, color=hx(p["text"]),
                        fontfamily=font, va="top", wrap=True)
                y -= step

    if kind not in ("title", "section", "quote"):
        band(5, 7.5, 90, 0.15, _mix(p["muted"], p["bg"], 0.72))
        if deck_title:
            ax.text(5, 5.4, deck_title, fontsize=8, color=hx(p["muted"]), fontfamily=font)
        ax.text(95, 5.4, f"{index} / {total}", fontsize=8, color=hx(p["muted"]),
                ha="right", fontfamily=font)

    out = str(path)
    fig.savefig(out, facecolor=fig.get_facecolor())
    plt.close(fig)
    return out


# =========================================================================== facade
def build(outdir: str | Path, spec: dict) -> dict:
    """Build whatever the spec asks for and return {files:[...], previews:[...]}."""
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    spec = ensure_spec(spec)
    style = spec["style"]
    made: list[str] = []
    previews: list[str] = []

    charts = {c.get("id"): c for c in (spec.get("charts") or []) if isinstance(c, dict)}
    for cid, c in charts.items():
        if c.get("id"):
            made.append(chart(outdir / f"{cid}.png", c, style))

    for b in spec.get("blocks", []):
        if isinstance(b, dict) and b.get("kind") == "chart" and b.get("chart_id") in charts:
            b.setdefault("image", str(outdir / f"{b['chart_id']}.png"))

    fmt = (spec.get("format") or "pptx").lower()
    if fmt in ("pptx", "presentation", "slides"):
        made.append(pptx_deck(outdir / "presentation.pptx", spec, style))
        slides = [{"kind": "title", "title": spec["title"], "subtitle": spec.get("subtitle", ""),
                   "author": spec.get("author", "")}] + \
                 [b for b in spec["blocks"] if isinstance(b, dict)] + \
                 [{"kind": "section", "title": spec.get("closing", "Спасибо!")}]
        for i, sl in enumerate(slides[:12], start=1):
            try:
                previews.append(slide_preview(outdir / f"slide_{i:02d}.png", sl, style,
                                              index=i, total=len(slides),
                                              deck_title=spec["title"]))
            except Exception:
                pass
    elif fmt in ("docx", "word", "document"):
        made.append(docx_report(outdir / "document.docx", spec, style))
    elif fmt in ("xlsx", "excel", "sheet"):
        made.append(xlsx_report(outdir / "workbook.xlsx", spec, style))
    elif fmt == "pdf":
        made.append(pdf_report(outdir / "document.pdf", spec, style))
    else:
        made.append(docx_report(outdir / "document.docx", spec, style))

    if spec.get("pdf_too"):
        try:
            made.append(pdf_report(outdir / "document.pdf", spec, style))
        except Exception:
            pass

    return {"files": made, "previews": previews, "style": style}
