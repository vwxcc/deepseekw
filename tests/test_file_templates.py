from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from sandbox.designs import templates


class FileTemplateTests(unittest.TestCase):
    def test_catalog_exposes_formats_templates_and_design_styles(self):
        catalog = templates.catalog()
        self.assertGreaterEqual(len(catalog["templates"]), 25)
        self.assertEqual(len(catalog["styles"]), 57)
        formats = {item["format"] for item in catalog["templates"]}
        self.assertTrue({"pptx", "docx", "xlsx", "svg", "png", "pdf"} <= formats)
        self.assertIn("presentation.pitch", {item["id"] for item in catalog["templates"]})

    def test_rejects_unknown_or_mismatched_templates(self):
        with self.assertRaisesRegex(ValueError, "Неизвестный шаблон"):
            templates.validate_plan({"format": "pptx", "template_id": "pptx.magic"})
        with self.assertRaisesRegex(ValueError, "предназначен для формата"):
            templates.validate_plan({"format": "docx", "template_id": "presentation.pitch"})

    def test_plan_is_data_not_executable_code_and_filename_is_sanitised(self):
        clean = templates.validate_plan({
            "format": "svg",
            "template_id": "vector.infographic",
            "title": "Схема / безопасный путь",
            "filename": "../../bad:name.svg",
            "blocks": [{"kind": "text", "text": "Только содержимое"}],
        })
        self.assertEqual(clean["filename"], "bad-name.svg")
        self.assertNotIn("..", clean["filename"])
        self.assertEqual(clean["blocks"][0]["kind"], "text")

    def test_builds_svg_and_png_from_the_same_declarative_plan(self):
        plan = {
            "format": "svg",
            "template_id": "vector.infographic",
            "title": "Сравнение вариантов",
            "subtitle": "Короткий аналитический обзор",
            "style": "editorial",
            "palette": {"background": "#F7F1E6", "primary": "#66594A", "accent": "#C36C45"},
            "blocks": [
                {"kind": "metrics", "title": "Главные цифры",
                 "items": [{"value": "42%", "label": "рост"}, {"value": "3,5", "label": "млн"}]},
                {"kind": "bullets", "title": "Выводы", "items": ["Первый", "Второй"]},
            ],
        }
        with tempfile.TemporaryDirectory() as tmp:
            svg_result = templates.build_from_plan(tmp, plan)
            svg_path = Path(svg_result["files"][0])
            self.assertTrue(svg_path.is_file())
            self.assertIn("<svg", svg_path.read_text(encoding="utf-8"))
            png_result = templates.build_from_plan(tmp, {**plan, "format": "png",
                "template_id": "image.infographic", "filename": "comparison.png"})
            png_path = Path(png_result["files"][0])
            self.assertTrue(png_path.is_file())
            from PIL import Image
            with Image.open(png_path) as image:
                self.assertEqual(image.format, "PNG")
                self.assertGreaterEqual(image.width, 900)

    def test_builds_office_files_from_template_blocks(self):
        with tempfile.TemporaryDirectory() as tmp:
            pptx = templates.build_from_plan(tmp, {
                "format": "pptx", "template_id": "presentation.pitch",
                "title": "План продукта", "style": "stripe",
                "filename": "product-plan.pptx",
                "slides": [
                    {"layout": "metrics", "title": "Результаты",
                     "items": [{"value": "42%", "label": "рост"}]},
                    {"layout": "bullets", "title": "Следующие шаги",
                     "items": ["Исследование", "Прототип"]},
                ],
            })
            pptx_file = Path(next(path for path in pptx["files"] if path.endswith(".pptx")))
            self.assertTrue(pptx_file.is_file())
            from pptx import Presentation
            self.assertGreaterEqual(len(Presentation(str(pptx_file)).slides), 4)

            docx = templates.build_from_plan(tmp, {
                "format": "docx", "template_id": "document.editorial",
                "title": "Краткий отчёт", "filename": "report.docx",
                "sections": [
                    {"kind": "h1", "text": "Выводы"},
                    {"kind": "bullets", "items": ["Первый вывод", "Второй вывод"]},
                    {"kind": "table", "head": ["Показатель", "Значение"], "rows": [["Рост", "42%"]]},
                ],
            })
            docx_file = Path(next(path for path in docx["files"] if path.endswith(".docx")))
            self.assertTrue(docx_file.is_file())
            from docx import Document
            self.assertGreaterEqual(len(Document(str(docx_file)).paragraphs), 2)

            xlsx = templates.build_from_plan(tmp, {
                "format": "xlsx", "template_id": "spreadsheet.analytics",
                "title": "Результаты", "filename": "metrics.xlsx",
                "sheets": [{"name": "Метрики", "head": ["Показатель", "Значение"],
                            "rows": [["Рост", 42], ["Конверсия", 3.7], ["Выручка", 1200]}],
            })
            xlsx_file = Path(next(path for path in xlsx["files"] if path.endswith(".xlsx")))
            self.assertTrue(xlsx_file.is_file())
            from openpyxl import load_workbook
            wb = load_workbook(xlsx_file, data_only=False)
            self.assertEqual(wb["Метрики"]["A1"].value, "Показатель")
            self.assertEqual(wb["Метрики"]["B2"].value, 42)


if __name__ == "__main__":
    unittest.main()
